"""Excel（xlsx）导出：15 列固定结构，openpyxl 生成。

排版约定：
- 第 1 行即表头（不设标题行），冻结 A2，A1:O{n} 自动筛选，全表自动换行。
- 配色：表头藏青（1F3864）+ 白字；编号列与隔行灰（D9D9D9 / F2F2F2）；细灰边框（BFBFBF）。
- 日期一律 YYYY-MM-DD 文本；原文链接为完整 URL；不使用任何合并单元格与公式。
- 公式注入防护：以 = + - @ 开头的原始文本会加前导单引号。
- demo payload 生成的工作簿在文档属性、工作表名与文件名上均明确标注 DEMO。
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

from .util import sanitize_cell_text

# 使用 8 位 ARGB（FF 不透明），避免部分阅读器把 6 位色值当成透明
HEADER_FILL_HEX = "FF1F3864"  # 藏青
BORDER_GREY_HEX = "FFBFBFBF"
BAND_GREY_HEX = "FFF2F2F2"
INDEX_GREY_HEX = "FFD9D9D9"
LINK_FONT_HEX = "FF1F4E79"
HEADER_FONT_HEX = "FFFFFFFF"

MAX_CELL_CHARS = 32000  # Excel 单元格上限 32767

COLUMNS: tuple[tuple[str, str, int], ...] = (
    ("编号", "index", 8),
    ("跟踪周次", "week", 13),
    ("监测日期", "monitor_date", 13),
    ("国家/地区", "jurisdiction", 12),
    ("发布机构/立法机关", "authority", 28),
    ("文件类型", "event_type", 16),
    ("文件名称", "title", 44),
    ("摘要", "summary", 56),
    ("生效/拟生效日期", "effective_date", 17),
    ("原文链接", "source_url", 48),
    ("监测渠道", "channel", 18),
    ("是否为新增/实质更新", "substantive_flag", 19),
    ("更新类别", "update_type", 16),
    ("与公司相关性初判", "relevance", 17),
    ("涉及主题", "topics", 26),
)

HEADERS = tuple(column[0] for column in COLUMNS)

RELEVANCE_LABELS = {
    "high": "高",
    "medium": "中",
    "low": "低",
    "irrelevant": "不相关",
    "needs_review": "待人工复核",
}

def relevance_label(value: Any) -> str:
    return RELEVANCE_LABELS.get(str(value or ""), "待人工复核")


def substantive_label(event: dict[str, Any]) -> str:
    """是否为新增/实质更新。信息不足时留空，不臆测。"""
    status = str(event.get("status") or "")
    update_type = str(event.get("update_type") or "")
    if status == "New" or update_type == "New":
        return "新增"
    flag = event.get("is_substantive")
    if flag is True:
        return "实质更新"
    if flag is False:
        return "非实质更新"
    return ""


def channel_label(event: dict[str, Any], names: dict[str, str]) -> str:
    source = str(event.get("source") or "")
    return names.get(source) or source


def build_rows(payload: dict[str, Any]) -> list[list[Any]]:
    events = payload.get("events") or []
    period = payload.get("period") or {}
    names = {str(s.get("id")): str(s.get("name")) for s in (payload.get("sources") or []) if isinstance(s, dict)}
    monitor_date = _date_only(payload.get("generated_at")) or str(period.get("end") or "")

    rows: list[list[Any]] = []
    for index, event in enumerate(events, start=1):
        if not isinstance(event, dict):
            continue
        rows.append(
            [
                index,
                period.get("week") or "",
                monitor_date,
                event.get("jurisdiction") or "Unknown",
                event.get("authority") or "",
                event.get("event_type") or "",
                event.get("title") or "",
                event.get("summary") or "",
                event.get("effective_date") or "",
                event.get("source_url") or "",
                channel_label(event, names),
                substantive_label(event),
                event.get("update_type") or "",
                relevance_label(event.get("relevance")),
                "、".join(str(t) for t in (event.get("topics") or [])),
            ]
        )
    return rows


def _date_only(value: Any) -> str:
    text = str(value or "")
    if len(text) >= 10 and text[4] == "-" and text[7] == "-":
        return text[:10]
    return ""


def build_workbook(payload: dict[str, Any]) -> Any:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover
        raise RuntimeError("openpyxl 未安装，无法导出 Excel") from exc

    mode = str(payload.get("mode") or "live")
    is_demo = mode == "demo"
    period = payload.get("period") or {}
    rows = build_rows(payload)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "DEMO-合规监测" if is_demo else "合规监测"

    header_fill = PatternFill("solid", fgColor=HEADER_FILL_HEX)
    band_fill = PatternFill("solid", fgColor=BAND_GREY_HEX)
    index_fill = PatternFill("solid", fgColor=INDEX_GREY_HEX)
    thin = Side(style="thin", color=BORDER_GREY_HEX)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    header_font = Font(bold=True, color=HEADER_FONT_HEX, size=11)
    link_font = Font(color=LINK_FONT_HEX, underline="single")
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    top_left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    top_center = Alignment(horizontal="center", vertical="top", wrap_text=True)

    for column_index, (header, _key, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=column_index, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center
        cell.border = border
        sheet.column_dimensions[get_column_letter(column_index)].width = width
    sheet.row_dimensions[1].height = 30

    for row_offset, row in enumerate(rows):
        excel_row = row_offset + 2
        banded = row_offset % 2 == 1
        for column_index, value in enumerate(row, start=1):
            header = HEADERS[column_index - 1]
            cell = sheet.cell(row=excel_row, column=column_index)
            if header == "编号":
                cell.value = int(value) if isinstance(value, int) else sanitize_cell_text(value)
            else:
                cell.value = _clip(sanitize_cell_text(value))
            cell.border = border
            if header == "编号":
                cell.fill = index_fill
                cell.alignment = top_center
            elif banded:
                cell.fill = band_fill
                cell.alignment = top_center if header in ("跟踪周次", "监测日期", "国家/地区") else top_left
            else:
                cell.alignment = top_center if header in ("跟踪周次", "监测日期", "国家/地区") else top_left
            if header == "原文链接" and isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink = cell.value
                cell.font = link_font
            if header in ("跟踪周次", "监测日期", "生效/拟生效日期"):
                cell.number_format = "@"

    last_row = max(1, len(rows) + 1)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{last_row}"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.sheet_properties.tabColor = HEADER_FILL_HEX

    _apply_metadata(workbook, payload, is_demo=is_demo, period=period, row_count=len(rows))
    return workbook


def _clip(value: str) -> str:
    if isinstance(value, str) and len(value) > MAX_CELL_CHARS:
        return value[: MAX_CELL_CHARS - 1] + "…"
    return value


def _apply_metadata(workbook: Any, payload: dict[str, Any], *, is_demo: bool, period: dict[str, Any], row_count: int) -> None:
    props = workbook.properties
    props.creator = "Compliance Radar"
    props.lastModifiedBy = "Compliance Radar"
    generated_at = payload.get("generated_at") or "（尚未生成）"
    week = period.get("week") or ""

    if is_demo:
        props.title = "AI 合规监测周报（DEMO 演示数据）"
        props.subject = "DEMO：全部内容为演示数据，不代表任何真实监测结果"
        props.category = "DEMO"
        props.keywords = "DEMO,演示数据,compliance-radar,demo-only"
        props.description = (
            "【演示数据 / DEMO】本工作簿由 public/data/demo.json 生成，"
            "所有条目均为示例，禁止用于任何合规判断或对外汇报。"
            f" 周期：{week}；条目数：{row_count}；生成时间：{generated_at}。"
        )
    else:
        props.title = "AI 合规监测周报"
        props.subject = f"监测周期 {period.get('start', '')} ~ {period.get('end', '')}（{week}）"
        props.category = "合规监测"
        props.keywords = "compliance,合规监测,regulatory"
        props.description = (
            f"真实监测数据。周期：{period.get('start', '')} ~ {period.get('end', '')}（{week}）；"
            f"条目数：{row_count}；生成时间：{generated_at}。"
            "“与公司相关性初判”为 AI 初筛结果，不构成法律意见；"
            "官方链接被取回也不等于事实核验完成，需人工复核后方可用于决策。"
        )


def export_to_path(payload: dict[str, Any], path: Path | str) -> Path:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook(payload)
    workbook.save(target)
    return target


def export_bytes(payload: dict[str, Any]) -> bytes:
    buffer = io.BytesIO()
    workbook = build_workbook(payload)
    workbook.save(buffer)
    return buffer.getvalue()


def report_filename(run_id: str, mode: str = "live") -> str:
    return f"{run_id}-demo.xlsx" if mode == "demo" else f"{run_id}.xlsx"


def column_count() -> int:
    return len(COLUMNS)


def column_headers() -> list[str]:
    return list(HEADERS)
